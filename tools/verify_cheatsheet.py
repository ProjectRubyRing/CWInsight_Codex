"""Validate artifact consistency and selected documented query semantics.

This is intentionally not an AWS Logs Insights parser or emulator. Regular
expressions are checked against Python's overlapping syntax, while account,
service availability and execution remain untested without an AWS session.
"""
from __future__ import annotations

from datetime import datetime, timedelta, timezone
import json
from pathlib import Path
import re
import warnings
import xml.etree.ElementTree as ET
from zipfile import ZipFile

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'outputs' / 'cloudwatch_insights_20261016'
NAME = 'CloudWatch_Logs_Insights_Cheatsheet'
checks: list[str] = []


def check(condition: bool, description: str) -> None:
    if not condition:
        raise AssertionError(description)
    checks.append(description)


def run() -> None:
    manifest = json.loads((OUT/'qa'/'manifest.json').read_text(encoding='utf-8'))
    queries = {q['id']:q for q in manifest['queries']}
    check(list(queries)==[f'Q{i:02}' for i in range(1,49)],'48 unique sequential query IDs')
    markdown = (OUT/f'{NAME}.md').read_text(encoding='utf-8')
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter('always')
        book = openpyxl.load_workbook(OUT/f'{NAME}.xlsx', data_only=False)
    check(not caught,'Excel readback without openpyxl warnings')
    check(len(book.sheetnames)==8,'8 populated worksheets')
    copy = book['01_コピー用']
    check(str(copy.freeze_panes)=='D6','Copy sheet freezes headers and identifying columns')
    check(copy.auto_filter.ref=='A5:E53','All 48 queries included in filter range')
    check(len(re.findall(r'<a id="q\d{2}"></a>',markdown))==48,'48 Markdown query anchors')
    source_ids = {s[0] for s in manifest['sources']}
    for q in queries.values():
        one_line = ' | '.join(q['lines'])
        check(copy[q['copyCell']].value==one_line,f'{q["id"]} Excel query matches canonical source')
        check(not re.search(r'[\r\n\t"]',one_line),f'{q["id"]} single-line clipboard text')
        check(not re.search(r'\bSOURCE\b',one_line),f'{q["id"]} no console-incompatible SOURCE')
        check(q['lines'][0].startswith(('fields ','filter ')),f'{q["id"]} begins with standalone command')
        check(all(q[k] for k in ['ready','change','purpose','analogy','sample','expected','background','caution','next']),f'{q["id"]} complete explanations')
        check(set(q['refs'])<=source_ids,f'{q["id"]} valid source IDs')
        block = markdown.split(f'<a id="{q["id"].lower()}"></a>',1)[1].split('<a id=',1)[0]
        check('\n| '.join(q['lines']) in block,f'{q["id"]} Markdown query matches canonical source')
        for field in ['purpose','analogy','sample','expected','background','caution','next']:
            check(q[field] in block,f'{q["id"]} Markdown {field} preserved')
        loc=manifest['locations'][q['id']]
        detail=book[loc['sheet']]
        check(detail.cell(loc['row'],2).value==q['id'],f'{q["id"]} detail anchor exists')
        check(detail.cell(loc['row']+3,3).value=='\n| '.join(q['lines']),f'{q["id"]} detail query matches')
        for field, offset in [('ready',1),('change',2),('purpose',4),('analogy',5),('sample',9),('expected',10),('background',11),('caution',12),('next',13)]:
            check(detail.cell(loc['row']+offset,3).value==q[field],f'{q["id"]} Excel {field} preserved')
        for n, step in enumerate(q['steps']):
            check(detail.cell(loc['row']+6+n,3).value==step,f'{q["id"]} Excel processing step {n+1} preserved')
    for sheet in book:
        check(sheet.sheet_view.showGridLines is False,f'{sheet.title} gridlines hidden')
        check(sheet.print_area is not None,f'{sheet.title} print area set')
        check(int(sheet.page_setup.paperSize)==8,f'{sheet.title} A3 paper configured')
        for row in sheet:
            for cell in row:
                if cell.value is None:
                    continue
                check(cell.font.name=='Meiryo UI',f'{sheet.title}!{cell.coordinate} Meiryo UI')
                check(cell.data_type!='f',f'{sheet.title}!{cell.coordinate} contains literal reference content')
                check(cell.data_type!='e',f'{sheet.title}!{cell.coordinate} no Excel error')
                check('\ufffd' not in str(cell.value),f'{sheet.title}!{cell.coordinate} no replacement characters')
                if cell.row >= 5:
                    check(cell.alignment.wrap_text is True,f'{sheet.title}!{cell.coordinate} wraps text')
                if cell.font.color and cell.font.color.type=='rgb':
                    color=cell.font.color.rgb[-6:]
                    check(color[:2]==color[2:4]==color[4:6],f'{sheet.title}!{cell.coordinate} monochrome text')
                if cell.fill.fgColor.type=='rgb':
                    color=cell.fill.fgColor.rgb[-6:]
                    check(color[:2]==color[2:4]==color[4:6],f'{sheet.title}!{cell.coordinate} monochrome fill')
                if cell.hyperlink and cell.hyperlink.location:
                    destination=cell.hyperlink.location
                    match=re.fullmatch(r"'([^']+)'!([A-Z]+[0-9]+)",destination)
                    check(match is not None and match[1] in book.sheetnames,f'{sheet.title}!{cell.coordinate} valid internal hyperlink')
                    check(book[match[1]][match[2]].value is not None,f'{sheet.title}!{cell.coordinate} hyperlink target populated')
        for dim in sheet.row_dimensions.values():
            check(not dim.height or dim.height<=409.5,f'{sheet.title} row height within Excel limit')

    def regex_for(qid: str, contains: str = '') -> re.Pattern:
        for line in queries[qid]['lines']:
            if contains not in line:
                continue
            m=re.search(r' /((?:\\.|[^/])*)/(?:$| )',line)
            if m:
                pattern=re.sub(r'\(\?<([A-Za-z][A-Za-z0-9]*)>',r'(?P<\1>',m[1])
                return re.compile(pattern)
        raise AssertionError(f'Missing regex in {qid}')

    for q in queries.values():
        for line in q['lines']:
            for m in re.finditer(r' /((?:\\.|[^/])*)/(?:$| )',line):
                pattern=re.sub(r'\(\?<([A-Za-z][A-Za-z0-9]*)>',r'(?P<\1>',m[1])
                re.compile(pattern)
                checks.append(f'{q["id"]} regex compiles in overlapping Python syntax')

    r=regex_for('Q02')
    for text,expected in [('ERROR failed',True),('error timeout',True),('FATAL stop',True),('SEVERE issue',True),('INFO errorCount=0',False),('NOERROR',False),('INFO previous ERROR cleared',True),('java.lang.OutOfMemoryError',False),('ERROR123',True)]:
        check(bool(r.search(text))==expected,f'Q02 boundary: {text}')
    check(bool(regex_for('Q03').search('INFO no error detected')),'Q03 documented false positive retained')
    candidates=['healthcheck OK','heartbeat OK','healthcheck heartbeat','payment failed']
    survivors=[s for s in candidates if 'healthcheck' not in s and 'heartbeat' not in s]
    check(survivors==['payment failed'],'Q10 AND exclusion truth table')
    check('healthcheck OK' in [s for s in candidates if 'healthcheck' not in s or 'heartbeat' not in s],'Q10 documented OR mistake demonstrated')
    r=regex_for('Q11')
    check([s for s in ['HealthCheck OK','HEARTBEAT OK','ERROR payment failed'] if not r.search(s)]==['ERROR payment failed'],'Q11 case-insensitive exclusions')
    jst=timezone(timedelta(hours=9))
    start=int(datetime(2026,10,1,9,tzinfo=jst).timestamp()*1000)
    end=int(datetime(2026,10,1,10,tzinfo=jst).timestamp()*1000)
    check((start,end)==(1790812800000,1790816400000),'Q15/Q16 independently calculated JST epoch values')
    check([x for x in [start-1,start,end-1,end] if start<=x<end]==[start,end-1],'Q15 start inclusive / end exclusive at millisecond boundaries')
    check(int(datetime(2026,10,1,9,5,tzinfo=jst).timestamp()*1000)==1790813100000,'Q39 five-minute end time')
    check(15*60*1000==900000 and 24*60*60*1000==86400000,'Q13/Q14 unit conversion')
    r=regex_for('Q18')
    check(r.search('111122223333:/example/eap/prod')['groupName']=='/example/eap/prod','Q18 account prefix removed')
    check(r.search('444455556666:/example/eap/prod')['groupName']=='/example/eap/prod','Q18 documented cross-account name collision')
    r=regex_for('Q22')
    check(bool(r.search('INFO WFLYSRV0025: started')),'Q22 includes normal startup code')
    check(not r.search('ARJUNA012117: example'),'Q22 prefix selection is not all EAP libraries')
    check(not r.search('WFLYCTL00130: example'),'Q22 code digit boundary')
    r=regex_for('Q28')
    check(r.search('WFLYCTL0013: first WFLYCTL0080: second')['eapCode']=='WFLYCTL0013','Q28 first code, not every code')
    for code in ['IJ000453','ARJUNA012117','UT005023','MSC000001','HHH000342']:
        check(r.search(code+': example')['eapCode']==code,f'Q28 extracts {code}')
    r=regex_for('Q30')
    for s in ['java.lang.NullPointerException','java.sql.SQLException','java.lang.OutOfMemoryError','java.lang.Error','java.lang.Exception']:
        check(bool(r.search(s)),f'Q30 qualified class {s}')
    check(not r.search('Exception occurred'),'Q30 unqualified class not included')
    chain='java.lang.RuntimeException: outer\nCaused by: java.sql.SQLException: middle\nCaused by: java.net.ConnectException: inner'
    check(regex_for('Q36').search(chain)['causeClass']=='java.sql.SQLException','Q36 first cause is not deepest cause')
    check(regex_for('Q42').search(chain)['exceptionClass']=='java.lang.RuntimeException','Q42 first outer exception counted')
    for qid in ['Q43']:
        cmds=queries[qid]['lines']
        d=next(i for i,v in enumerate(cmds) if v.startswith('dedup '))
        check(all(v.startswith('limit ') for v in cmds[d+1:]),'dedup followed only by limit')
    check(queries['Q45']['lines'][-1]=='limit any 20' and not any(x.startswith('sort') for x in queries['Q45']['lines']),'Q45 unordered early-exit query')
    check(queries['Q46']['lines'][-1]=='estimate','Q46 estimate is last command')
    check('bin(300s)' not in ' '.join(queries['Q40']['lines']),'Q40 avoids capped seconds unit')
    names=['/example/eap/prod','/example/eap/test','/archive/example/eap/prod','/example/eap-old']
    check([n for n in names if n.startswith('/example/eap/')]==names[:2],'Q48 prefix sample')
    renderlog=json.loads((OUT/'qa'/'render_log.json').read_text(encoding='utf-8'))
    check(all(r['result']=='rendered' for r in renderlog),'All requested sheet/range renders generated')
    with ZipFile(OUT/f'{NAME}.xlsx') as z:
        check(z.testzip() is None,'XLSX ZIP integrity')
        for name in z.namelist():
            if name.endswith(('.xml','.rels')):
                ET.fromstring(z.read(name))
        checks.append('All XLSX XML parts parse successfully')
    for file in [ROOT/'README.md',* (ROOT/'tools').glob('*.mjs'),*(ROOT/'tools').glob('*.py'),OUT/f'{NAME}.md']:
        raw=file.read_bytes()
        check(not raw.startswith(b'\xef\xbb\xbf'),f'{file.name} UTF-8 without BOM')
        text=raw.decode('utf-8')
        check(not re.search(r'[ \t]+$',text,re.MULTILINE),f'{file.name} no trailing whitespace')
        check('\r' not in text,f'{file.name} consistent LF line endings')
    report={'result':'PASS','checks':len(checks),'query_count':len(queries),'sheet_count':len(book.sheetnames),'source_count':len(source_ids),'cloudwatch_execution':'NOT RUN','regex_engine':'Python subset compatibility, not AWS engine','native_excel_visual_check':'Not performed by this script','checks_detail':checks}
    (OUT/'qa'/'verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({k:v for k,v in report.items() if k!='checks_detail'},ensure_ascii=True,indent=2))


if __name__=='__main__':
    run()
